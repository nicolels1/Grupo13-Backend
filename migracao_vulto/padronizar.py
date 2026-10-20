# Padroniza os dados da Vulto numa área de staging, antes da migração para o banco da Casa Lorenzi.
# Aplica só as transformações automáticas (data, número, CPF, e-mail, telefone, texto, encoding e
# cabeçalho de planilha). O que depende de decisão — canal, status, categoria, pontos, cor, SKU
# canônico, total do app em centavos, datas da planilha manual da Batel — é copiado como veio para a
# próxima etapa tratar.
# Não altera os arquivos da Vulto (o SQLite é aberto só para leitura) nem grava no banco da Lorenzi:
# escreve um CSV por tabela/aba em --destino, mais pendencias.csv com cada valor que não deu para
# padronizar (fonte, linha, coluna, valor original e motivo).
# Uso, na raiz do repositório com o venv ativo (a origem vem de VULTO_DATA_DIR no .env ou de --origem):
#   pip install -r migracao_vulto/requirements.txt
#   python -m migracao_vulto.padronizar --destino staging_vulto
# O destino tem CPF, e-mail e telefone de clientes: fica fora do git (.gitignore) e não deve ser compartilhado.
import argparse
import re
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from migracao_vulto import normalizadores as n
from migracao_vulto.normalizadores import ValorInvalido
from src.config.settings import VULTO_DATA_DIR

ARQUIVO_BANCO = "vulto.db"
ARQUIVO_CLIENTES = "clientes_fidelidade_FINAL_v3 (1).xlsx"
ARQUIVO_CATALOGO = "catalogo_vulto_2026_REVISADO_final.xlsx"
ARQUIVO_CONTAGEM = "contagem_inventario_set26.xlsx"
ARQUIVO_BATEL = "vendas_batel_MANUAL_marco.xlsx"
ARQUIVO_ATENDIMENTO = "atendimentos_export_whatsapp_insta.csv"


# ---------- pendências ----------

class Pendencias:
    """Junta os valores que não deu para padronizar; cada um vira uma linha de pendencias.csv."""

    def __init__(self):
        self.linhas = []

    def registrar(self, fonte: str, linha, coluna: str, valor, motivo: str):
        self.linhas.append(dict(fonte=fonte, linha=linha, coluna=coluna, valor=str(valor), motivo=motivo))

    # aplica a função ao valor; se ela levantar ValorInvalido, registra a pendência e devolve `padrao`
    def aplicar(self, funcao, valor, fonte: str, linha, coluna: str, padrao=None):
        try:
            return funcao(valor)
        except ValorInvalido as erro:
            self.registrar(fonte, linha, coluna, valor, str(erro))
            return padrao


# ---------- utilitários de DataFrame ----------

# Aplica uma função de normalizadores a cada coluna listada em `regras` ({coluna: função}).
# A coluna linha_origem diz de onde veio cada registro (id no banco ou linha no Excel), para rastrear
# a pendência e, depois, a migração.
def aplicar_regras(df: pd.DataFrame, regras: dict, fonte: str, pend: Pendencias) -> pd.DataFrame:
    df = df.copy()
    for coluna, funcao in regras.items():
        df[coluna] = [pend.aplicar(funcao, valor, fonte, linha, coluna)
                      for linha, valor in zip(df["linha_origem"], df[coluna])]
    return df


# Aplica uma função que devolve um par (como contato e pedido) e separa em duas colunas novas
def aplicar_par(df: pd.DataFrame, coluna: str, funcao, novas: tuple[str, str], fonte: str,
                pend: Pendencias) -> pd.DataFrame:
    df = df.copy()
    pares = [pend.aplicar(funcao, valor, fonte, linha, coluna, padrao=(None, None))
             for linha, valor in zip(df["linha_origem"], df[coluna])]
    df[novas[0]] = [p[0] for p in pares]
    df[novas[1]] = [p[1] for p in pares]
    return df


# Tira espaço das pontas de todo texto e troca "", "-", "N/A"... por vazio. É o mínimo seguro para
# qualquer coluna, inclusive as que ficam como vieram ("loja " vira "loja", o de-para vem depois)
def limpar_textos(df: pd.DataFrame) -> pd.DataFrame:
    return df.map(lambda v: (None if n.vazio(v) else v.strip()) if isinstance(v, str) else v)


# Lê uma aba de planilha sem deixar o pandas converter nada (dtype=object mantém número, texto e data
# como estão, que é o que as funções de normalização esperam) e numera a linha como o Excel mostra
def ler_aba(arquivo: Path, aba: str, linha_cabecalho: int = 0) -> pd.DataFrame:
    df = pd.read_excel(arquivo, sheet_name=aba, header=linha_cabecalho, dtype=object)
    df = df.dropna(how="all")  # linhas totalmente vazias no meio ou no fim da aba
    df.insert(0, "linha_origem", [f"{aba}!{i + linha_cabecalho + 2}" for i in df.index])
    return limpar_textos(df)


def gravar(df: pd.DataFrame, destino: Path, nome: str):
    caminho = destino / f"{nome}.csv"
    df.to_csv(caminho, index=False, encoding="utf-8")
    print(f"  {caminho.name}: {len(df)} linhas")


# ---------- 1. banco vulto.db ----------

# Lê a tabela inteira como bytes e decodifica cada texto (UTF-8 ou latin-1), resolvendo o "S\xe3o Paulo"
def ler_tabela(conexao: sqlite3.Connection, tabela: str, chave: str = "id") -> pd.DataFrame:
    cursor = conexao.execute(f"select * from {tabela}")
    colunas = [c[0] for c in cursor.description]
    linhas = [[n.decodificar(v) for v in linha] for linha in cursor.fetchall()]
    df = pd.DataFrame(linhas, columns=colunas, dtype=object)
    if chave == "id":
        origem = [f"{tabela}#{i}" for i in df["id"]]
    else:  # estoque não tem id: a chave é (sku, loja_id)
        origem = [f"{tabela}#{s}/{l}" for s, l in zip(df["sku"], df["loja_id"])]
    df.insert(0, "linha_origem", origem)
    return limpar_textos(df)


def padronizar_banco(origem: Path, destino: Path, pend: Pendencias):
    print("vulto.db")
    # mode=ro: o arquivo da Vulto é a cópia congelada da migração e não pode ser alterado
    conexao = sqlite3.connect(f"file:{origem / ARQUIVO_BANCO}?mode=ro", uri=True)
    conexao.text_factory = bytes
    fonte = ARQUIVO_BANCO

    # lojas: só encoding e espaços (a cidade vinha em latin-1)
    gravar(ler_tabela(conexao, "lojas"), destino, "banco_lojas")

    # funcionarios: a senha em texto puro NÃO é copiada para o staging; nome com encoding corrigido
    funcionarios = ler_tabela(conexao, "funcionarios").drop(columns=["senha"])
    gravar(aplicar_regras(funcionarios, {"nome": n.texto}, fonte, pend), destino, "banco_funcionarios")

    # vendas: data em UTC, valores em Decimal e o cliente_ref separado em tipo + valor padronizado.
    # canal, status e obs ficam como vieram (de-para é padrão assumido); total do app em centavos também
    vendas = ler_tabela(conexao, "vendas")
    vendas = aplicar_regras(vendas, {"criado_em": n.data_hora, "total": n.decimal_br,
                                     "desconto": n.decimal_br, "frete": n.decimal_br}, fonte, pend)
    vendas = aplicar_par(vendas, "cliente_ref", n.contato, ("cliente_tipo", "cliente_valor"), fonte, pend)
    gravar(vendas, destino, "banco_vendas")

    # itens_venda: preço unitário em Decimal e SKU sem espaço/minúscula
    itens = ler_tabela(conexao, "itens_venda")
    itens = aplicar_regras(itens, {"preco_unit": n.decimal_br, "sku": n.sku, "descricao": n.texto}, fonte, pend)
    gravar(itens, destino, "banco_itens_venda")

    # compras (lado fidelidade da venda): mesmas regras de valor, data e cliente
    compras = ler_tabela(conexao, "compras")
    compras = aplicar_regras(compras, {"criado_em": n.data_hora, "valor": n.decimal_br}, fonte, pend)
    compras = aplicar_par(compras, "cliente_ref", n.contato, ("cliente_tipo", "cliente_valor"), fonte, pend)
    gravar(compras, destino, "banco_compras")

    # pontos_mov: CPF com os zeros à esquerda de volta e data em UTC. tipo e qtd ficam como vieram
    pontos = ler_tabela(conexao, "pontos_mov")
    gravar(aplicar_regras(pontos, {"cpf": n.cpf, "criado_em": n.data_hora}, fonte, pend), destino, "banco_pontos_mov")

    # entradas_estoque: data (3 formatos, dia primeiro), custo em Decimal e SKU. fornecedor fica como veio
    entradas = ler_tabela(conexao, "entradas_estoque")
    entradas = aplicar_regras(entradas, {"data": n.data, "custo_unit": n.decimal_br, "sku": n.sku}, fonte, pend)
    gravar(entradas, destino, "banco_entradas_estoque")

    # estoque: só a foto atual; data de atualização em UTC e SKU. Saldo negativo fica como veio
    estoque = ler_tabela(conexao, "estoque", chave="sku_loja")
    estoque = aplicar_regras(estoque, {"atualizado_em": n.data_hora, "sku": n.sku}, fonte, pend)
    gravar(estoque, destino, "banco_estoque")

    conexao.close()


# ---------- 2. planilha de clientes Vulto+ ----------

def padronizar_clientes(origem: Path, destino: Path, pend: Pendencias):
    print(ARQUIVO_CLIENTES)
    arquivo = origem / ARQUIVO_CLIENTES
    fonte = ARQUIVO_CLIENTES

    # Clientes: CPF, nome, e-mail e telefone. Data de cadastro em texto, nº de série do Excel, data ou
    # só mês/ano ("ago/26", marcado em data_cadastro_so_mes). nome_busca (sem acento, minúsculo) serve
    # para achar duplicados depois.
    # Pontos, categoria, aceita_msg, loja_cadastro e Obs ficam como vieram: são padrão assumido ou decisão
    clientes = ler_aba(arquivo, "Clientes")
    clientes = aplicar_regras(clientes, {"cpf": n.cpf, "E-mail": n.email, "Tel": n.telefone}, fonte, pend)
    clientes = aplicar_par(clientes, "Data Cadastro", n.data_ou_mes, ("Data Cadastro", "data_cadastro_so_mes"),
                           fonte, pend)
    clientes["nome_busca"] = clientes["Nome"].map(n.chave_busca)
    clientes["Nome"] = clientes["Nome"].map(n.nome_proprio)
    gravar(clientes, destino, "clientes_clientes")

    # clientes antigos (programa de 2023): CPF, nome, data da última compra (DD/MM/AA) e pontos inteiros
    antigos = ler_aba(arquivo, "clientes antigos (NAO APAGAR)")
    antigos = aplicar_regras(antigos, {"CPF/CNPJ": n.cpf, "ULT_COMPRA": n.data, "PONTOS_2023": n.inteiro},
                             fonte, pend)
    antigos["NOME COMPLETO"] = antigos["NOME COMPLETO"].map(n.nome_proprio)
    gravar(antigos, destino, "clientes_antigos")

    # Resgates: data (inclui "set/23", marcado em data_so_mes), pontos inteiros e desconto em Decimal.
    # A coluna cliente tem CPF ou só o nome: CPF vai para cliente_cpf, o resto fica em cliente_nome
    resgates = ler_aba(arquivo, "Resgates")
    resgates = aplicar_par(resgates, "data", n.data_ou_mes, ("data", "data_so_mes"), fonte, pend)
    resgates = aplicar_regras(resgates, {"pontos usados": n.pontos_milhar, "desconto R$": n.decimal_br}, fonte, pend)
    resgates["cliente_cpf"] = [_cpf_ou_nada(v) for v in resgates["cliente"]]
    resgates["cliente_nome"] = [None if c else n.nome_proprio(v)
                                for v, c in zip(resgates["cliente"], resgates["cliente_cpf"])]
    gravar(resgates, destino, "clientes_resgates")


# na aba Resgates, "cliente" que não é CPF é nome de pessoa (não é pendência, só outro tipo de dado)
def _cpf_ou_nada(valor):
    try:
        return n.cpf(valor)
    except ValorInvalido:
        return None


# ---------- 3. catálogo ----------

# Cada aba foi montada por uma pessoa e tem cabeçalho e nomes de coluna próprios (mesmo mapa do
# catalogo.py da Vulto). A aba "NAO USAR" entra porque tem SKUs de 2024 que ainda vendem; a aba
# de regras não tem produto. Tudo sai com as mesmas colunas e a aba de origem.
ABAS_CATALOGO = {
    "VERAO 25": dict(cabecalho=0, sku="SKU", nome="Produto", categoria="Categoria", cor="Cor", tamanho="Tam",
                     preco="Preço", custo="Custo", ean="EAN"),
    "inverno25": dict(cabecalho=2, sku="cod", nome="descricao", categoria="cat", cor="cor", tamanho="tamanho",
                      preco="preco venda", custo="custo unit", ean="ean13"),
    "Verão 26": dict(cabecalho=0, sku="Código", nome="Nome", categoria="Categoria", cor="Cor", tamanho="Tamanho",
                     preco="Preço (R$)", custo="Custo", ean="EAN"),
    "Atemporais": dict(cabecalho=0, sku="sku", nome="nome", categoria="categoria", cor="cor", tamanho="tam",
                       preco="preco", custo="custo_medio", ean="ean"),
    "PLANILHA ANTIGA NAO USAR": dict(cabecalho=1, sku="SKU", nome="DESCRICAO", categoria="GRUPO", cor="COR",
                                     tamanho="TAM", preco="PRECO", custo="CUSTO", ean="EAN"),
}
COLUNAS_CATALOGO = ["sku", "nome", "categoria", "cor", "tamanho", "preco", "custo", "ean"]


# Cores de fundo que as abas usam como marcação (legendas "vermelho = descontinuado" e
# "amarelo = revisar preço"); há dois tons de vermelho, conforme quem marcou
VERMELHOS = {"FFFF0000", "FFC00000"}
AMARELO = "FFFFFF00"


def _cor_fundo(celula) -> str | None:
    if celula.fill is None or not celula.fill.fill_type:
        return None
    return celula.fill.fgColor.rgb


# Lê as marcações que só existem na formatação (o pandas não enxerga): descontinuado é texto tachado ou
# a linha inteira em vermelho (SKU e nome); revisar_preco é o fundo amarelo. Só o SKU em vermelho, com
# o resto da linha normal, não diz o que é: fica sem valor e vira pendência.
def marcar_formatacao(df: pd.DataFrame, planilha, mapa: dict, aba: str, pend: Pendencias) -> pd.DataFrame:
    ws = planilha[aba]
    cabecalho = [c.value for c in ws[mapa["cabecalho"] + 1]]
    coluna_sku = cabecalho.index(mapa["sku"]) + 1
    coluna_nome = cabecalho.index(mapa["nome"]) + 1
    descontinuado, revisar = [], []
    for linha_origem in df["linha_origem"]:
        linha = int(linha_origem.rsplit("!", 1)[1])
        celula_sku, celula_nome = ws.cell(linha, coluna_sku), ws.cell(linha, coluna_nome)
        sku_vermelho = _cor_fundo(celula_sku) in VERMELHOS
        nome_vermelho = _cor_fundo(celula_nome) in VERMELHOS
        if celula_sku.font.strike or (sku_vermelho and nome_vermelho):
            descontinuado.append(True)
        elif sku_vermelho:
            descontinuado.append(None)
            pend.registrar(ARQUIVO_CATALOGO, linha_origem, "descontinuado", "só o SKU em vermelho",
                           "marcação parcial: não dá para saber se é descontinuado")
        else:
            descontinuado.append(False)
        revisar.append(_cor_fundo(celula_sku) == AMARELO)
    df = df.copy()
    df["descontinuado"] = descontinuado
    df["revisar_preco"] = revisar
    return df


# Coleção de cada aba. A Marina pôs o Inverno 26 dentro da aba Verão 26, abaixo de uma linha de título
# ("INVERNO 26 (PROVISÓRIO - MARINA) - AINDA SEM ABA PROPRIA"): o título sai e marca a coleção das
# linhas de baixo (C15, C9)
COLECAO_DA_ABA = {"VERAO 25": "Verão 25", "inverno25": "Inverno 25", "Verão 26": "Verão 26",
                  "Atemporais": "Atemporais", "PLANILHA ANTIGA NAO USAR": "2024 (aba antiga)"}
TITULO_COLECAO = re.compile(r"^\s*(INVERNO|VER[AÃ]O)\s+(\d{2})\b", re.IGNORECASE)


def marcar_colecao(df: pd.DataFrame, aba: str) -> pd.DataFrame:
    colecao, colecoes, titulos = COLECAO_DA_ABA[aba], [], []
    for indice, sku, nome, preco in zip(df.index, df["sku"], df["nome"], df["preco"]):
        m = TITULO_COLECAO.match(str(sku))
        if m and pd.isna(nome) and pd.isna(preco):  # linha de título, não de produto
            estacao = "Verão" if m.group(1).upper() != "INVERNO" else "Inverno"
            colecao = f"{estacao} {m.group(2)}"
            titulos.append(indice)
        colecoes.append(colecao)
    df = df.assign(colecao=colecoes)
    return df.drop(index=titulos)


def padronizar_catalogo(origem: Path, destino: Path, pend: Pendencias):
    print(ARQUIVO_CATALOGO)
    arquivo = origem / ARQUIVO_CATALOGO
    planilha = load_workbook(arquivo)  # com formatação, para as marcações de cor e tachado
    partes = []
    for aba, mapa in ABAS_CATALOGO.items():
        df = ler_aba(arquivo, aba, mapa["cabecalho"])
        # renomeia para os nomes comuns e descarta colunas de anotação (legenda de cor, obs, fórmula)
        df = df.rename(columns={mapa[c]: c for c in COLUNAS_CATALOGO})[["linha_origem", *COLUNAS_CATALOGO]]
        df = df[df["sku"].notna()]  # linha sem SKU é título, legenda ou espaço
        # cabeçalho repetido no meio da aba (linha copiada e colada junto com os produtos)
        df = df[df["sku"].astype(str).str.strip().str.lower() != mapa["sku"].lower()]
        df = marcar_colecao(df, aba)
        df = marcar_formatacao(df, planilha, mapa, aba, pend)
        df.insert(1, "aba", aba)
        df.insert(2, "colecao", df.pop("colecao"))
        partes.append(df)
    catalogo = pd.concat(partes, ignore_index=True)
    # preço "R$ 129,90" ou número, custo, EAN como texto, tamanho e SKU em maiúsculas, nome sem espaços.
    # categoria e cor ficam como vieram: o de-para é padrão assumido. Peça repetida entre abas também
    # fica (qual preço vale é decisão: coleção mais nova × menor preço)
    catalogo = aplicar_regras(catalogo, {"sku": n.sku, "nome": n.texto, "tamanho": n.tamanho,
                                         "preco": n.decimal_br, "custo": n.decimal_br, "ean": n.ean13},
                              ARQUIVO_CATALOGO, pend)
    gravar(catalogo, destino, "catalogo_produtos")


# ---------- 4. contagem física de set/26 ----------

# Mesma ideia do catálogo: cada loja contou com um layout. nivel diz se a linha é por SKU ou por
# modelo (a Savassi contou só por modelo, sem cor e tamanho, e isso não dá para desagregar).
# data_contagem só onde a aba diz a data; nas outras fica vazia.
ABAS_CONTAGEM = {
    "PINHEIROS 20-09": dict(cabecalho=2, codigo="SKU", quantidade="QTD", obs="OBS", nivel="sku",
                            unidade="Pinheiros", data="2026-09-20"),
    "Eldorado": dict(cabecalho=0, codigo="codigo", quantidade="qtde", obs="conferente", nivel="sku",
                     unidade="Eldorado", data=None),
    "SAVASSI (por modelo)": dict(cabecalho=0, codigo="MODELO", quantidade="QTD TOTAL", obs="DESCRIÇÃO",
                                 nivel="modelo", unidade="Savassi", data=None),
    "batel": dict(cabecalho=0, codigo="sku", quantidade="quantidade", obs="obs", nivel="sku",
                  unidade="Batel", data=None),
    "CD - posicao WMS 28-09": dict(cabecalho=0, codigo="SKU", quantidade="SALDO", reservado="RESERVADO",
                                   nivel="sku", unidade="CD Cajamar", data="2026-09-28"),
}


def padronizar_contagem(origem: Path, destino: Path, pend: Pendencias):
    print(ARQUIVO_CONTAGEM)
    arquivo = origem / ARQUIVO_CONTAGEM
    partes = []
    for aba, mapa in ABAS_CONTAGEM.items():
        df = ler_aba(arquivo, aba, mapa["cabecalho"])
        saida = pd.DataFrame({
            "linha_origem": df["linha_origem"],
            "unidade": mapa["unidade"],
            "data_contagem": mapa["data"],
            "nivel": mapa["nivel"],
            "codigo": df[mapa["codigo"]],
            "quantidade": df[mapa["quantidade"]],
            "reservado": df[mapa["reservado"]] if "reservado" in mapa else None,
            "observacao": df[mapa["obs"]] if "obs" in mapa else None,
        })
        partes.append(saida[saida["codigo"].notna()])
    contagem = pd.concat(partes, ignore_index=True)
    # código em maiúsculas e quantidades inteiras ("2 (1 no estoque + 1 vitrine)" vira 2 se as parcelas
    # somam o total); conversão para o SKU canônico fica para depois
    contagem = aplicar_regras(contagem, {"codigo": n.sku, "quantidade": n.quantidade_contada,
                                         "reservado": n.inteiro}, ARQUIVO_CONTAGEM, pend)
    gravar(contagem, destino, "contagem_set26")


# ---------- 5. vendas manuais da Batel (mar/26) ----------

def padronizar_batel(origem: Path, destino: Path, pend: Pendencias):
    print(ARQUIVO_BATEL)
    vendas = ler_aba(origem / ARQUIVO_BATEL, "MARÇO", linha_cabecalho=2)
    # CPF, valor e quantidade são automáticos. DATA NÃO é convertida: o Excel já leu "02/03" como
    # 3 de fevereiro em algumas linhas e outras não têm ano; fica o texto original para revisão
    vendas = aplicar_regras(vendas, {"CPF": n.cpf, "VALOR": n.decimal_br, "QTD": n.inteiro,
                                     "CLIENTE": n.texto, "PEÇAS": n.texto}, ARQUIVO_BATEL, pend)
    vendas["DATA"] = vendas["DATA"].map(lambda v: None if n.vazio(v) else str(v))
    gravar(vendas, destino, "batel_vendas_manuais")


# ---------- 6. export do atendimento (WhatsApp / Instagram / e-mail) ----------

def padronizar_atendimento(origem: Path, destino: Path, pend: Pendencias):
    print(ARQUIVO_ATENDIMENTO)
    fonte = ARQUIVO_ATENDIMENTO
    # latin-1, primeira linha "sep=;" e mensagens com quebra de linha entre aspas (o pandas já trata)
    df = pd.read_csv(origem / ARQUIVO_ATENDIMENTO, sep=";", encoding="latin-1", skiprows=1,
                     dtype=str, keep_default_na=False)
    df.insert(0, "linha_origem", [f"csv!{i + 3}" for i in df.index])  # +3: linha sep=; e cabeçalho
    df = limpar_textos(df)

    # id repetido com a linha inteira igual é a mesma exportação duas vezes: sai. Se o resto da linha
    # muda, não dá para saber qual vale: as duas ficam e viram pendência
    colunas_dado = [c for c in df.columns if c != "linha_origem"]
    df = df.drop_duplicates(subset=colunas_dado)
    for linha, valor in zip(df["linha_origem"], df["id"]):
        if (df["id"] == valor).sum() > 1:
            pend.registrar(fonte, linha, "id", valor, "id repetido com conteúdo diferente")

    # data em UTC (DD/MM/AAAA HH:MM em Brasília ou ISO com fuso), mensagem sem espaço nas pontas
    df = aplicar_regras(df, {"data_hora": n.data_hora, "mensagem": n.texto_longo, "nome_contato": n.texto},
                        fonte, pend)
    # contato: telefone (+55...), @ do Instagram ou e-mail; aqui não existe CPF
    df = aplicar_par(df, "contato", lambda v: n.contato(v, aceita_cpf=False), ("contato_tipo", "contato_valor"),
                     fonte, pend)
    # pedido citado: só classifica o formato (site, app, número solto); o vínculo é a próxima etapa
    df = aplicar_par(df, "pedido", n.pedido, ("pedido_tipo", "pedido_codigo"), fonte, pend)
    # tags numa lista separada por ";" e em minúsculas. canal, status e atendente ficam como vieram
    df["tags"] = df["tags"].map(lambda v: ";".join(n.tags(v)) or None)
    gravar(df, destino, "atendimentos")


# ---------- execução ----------

def main():
    parser = argparse.ArgumentParser(description="Padroniza os dados da Vulto (só as regras automáticas).")
    parser.add_argument("--origem", type=Path, default=VULTO_DATA_DIR,
                        help="pasta data/ do vulto-legado-back (padrão: VULTO_DATA_DIR do .env)")
    parser.add_argument("--destino", type=Path, default=Path("staging_vulto"), help="pasta dos CSVs padronizados")
    args = parser.parse_args()
    if args.origem is None:
        parser.error("informe --origem ou defina VULTO_DATA_DIR no .env")
    args.destino.mkdir(parents=True, exist_ok=True)

    # cada fonte é independente: se uma falhar, as outras já gravadas continuam valendo
    pend = Pendencias()
    for etapa in (padronizar_banco, padronizar_clientes, padronizar_catalogo, padronizar_contagem,
                  padronizar_batel, padronizar_atendimento):
        etapa(args.origem, args.destino, pend)

    # resumo das pendências por fonte e motivo, para saber o tamanho do que sobra para decidir
    pendencias = pd.DataFrame(pend.linhas, columns=["fonte", "linha", "coluna", "valor", "motivo"])
    gravar(pendencias, args.destino, "pendencias")
    if len(pendencias):
        print(pendencias.groupby(["fonte", "coluna", "motivo"]).size().to_string())


if __name__ == "__main__":
    main()
